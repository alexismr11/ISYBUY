"""Restriction au périmètre client, classeur Excel, synthèse JSON.

Rien ici ne lit un export ni ne modifie iSYBUY : ce module transforme le
résultat de mdf.detection.analyser() en livrables (classeur + JSON), après
avoir, le cas échéant, restreint le résultat au périmètre d'un client.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from mdf.modele import (
    LIBELLES_CAUSES,
    LIBELLES_SIGNAUX,
    Fournisseur,
    Palier,
    Population,
)

COULEUR_ENTETE = "0A8280"
COLONNES_SUIVI = ["Statut", "Opérateur", "Date traitement", "Commentaire"]


class ClientInconnu(Exception):
    pass


class ClientAmbigu(Exception):
    def __init__(self, correspondances: List[str]):
        self.correspondances = correspondances
        super().__init__(f"Plusieurs clients correspondent : {', '.join(correspondances)}")


def resoudre_client(groupes: Dict[str, str], nom_partiel: str) -> str:
    """Résout un nom de client (partiel, insensible à la casse) en identifiant de groupe."""
    cible = nom_partiel.strip().casefold()
    exact = [gid for gid, nom in groupes.items() if nom.casefold() == cible]
    if exact:
        return exact[0]
    correspondances = {gid: nom for gid, nom in groupes.items() if cible in nom.casefold()}
    if not correspondances:
        raise ClientInconnu(f"Client inconnu : « {nom_partiel} ». Voir --liste-clients.")
    if len(correspondances) > 1:
        raise ClientAmbigu(sorted(correspondances.values()))
    return next(iter(correspondances))


def _membres_prives_hors_client(membres: List[Fournisseur], id_client: str) -> set:
    return {m.id_groupe for m in membres if not m.est_public and m.id_groupe != id_client}


def _en_perimetre(membres: List[Fournisseur], id_client: str) -> bool:
    """Une fiche d'un autre client dans le groupe retire tout le groupe du périmètre.

    Fidèle au principe documenté : « la famille est retirée du périmètre
    plutôt que tronquée » — jamais de fuite, même partielle, vers un client.
    """
    concerne_client = any((not m.est_public and m.id_groupe == id_client) for m in membres) or any(
        m.est_public for m in membres
    )
    return concerne_client and not _membres_prives_hors_client(membres, id_client)


def restreindre_perimetre(analyse: dict, id_client: Optional[str], base_complete: bool) -> dict:
    if base_complete or id_client is None:
        return analyse

    return {
        "familles": [f for f in analyse["familles"] if _en_perimetre(f.membres, id_client)],
        "restructurations": [g for g in analyse["restructurations"] if _en_perimetre(g.membres, id_client)],
        "candidats_rattrapage": [c for c in analyse["candidats_rattrapage"] if _en_perimetre(c.membres, id_client)],
        "quarantaine": [
            q for q in analyse["quarantaine"] if q.fournisseur.est_public or q.fournisseur.id_groupe == id_client
        ],
    }


def _clients_concernes(membres: List[Fournisseur]) -> str:
    noms = sorted({m.nom_groupe or "Référentiel partagé" for m in membres})
    return ", ".join(noms)


def _style_entete(ws: Worksheet, nb_colonnes: int) -> None:
    fill = PatternFill(start_color=COULEUR_ENTETE, end_color=COULEUR_ENTETE, fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, nb_colonnes + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _feuille(wb: Workbook, titre: str, entetes: List[str]) -> Worksheet:
    ws = wb.create_sheet(titre)
    ws.append(entetes)
    _style_entete(ws, len(entetes))
    return ws


def _ajuster_largeurs(ws: Worksheet, largeurs: List[int]) -> None:
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largeur


def construire_classeur(analyse: dict, chemin: Path, contexte: dict) -> Path:
    """Écrit le classeur d'actions à sept onglets. Ne fusionne rien, n'appelle rien."""
    wb = Workbook()
    wb.remove(wb.active)

    familles = analyse["familles"]
    certaines = [f for f in familles if f.population == Population.DOUBLON_ACTIONNABLE and f.palier == Palier.CERTAINE]
    a_controler = [f for f in familles if f.population == Population.DOUBLON_ACTIONNABLE and f.palier == Palier.A_CONTROLER]
    arbitrages = [f for f in familles if f.population == Population.ARBITRAGE_INTER_TENANT]

    _feuille_synthese(wb, analyse, certaines, a_controler, arbitrages, contexte)
    _feuille_fusions(wb, "Fusions certaines", certaines, avec_controles=False)
    _feuille_fusions(wb, "Fusions à contrôler", a_controler, avec_controles=True)
    _feuille_arbitrages(wb, arbitrages)
    _feuille_restructurations(wb, analyse["restructurations"])
    _feuille_candidats(wb, analyse["candidats_rattrapage"])
    _feuille_quarantaine(wb, analyse["quarantaine"])

    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)
    return chemin


def _feuille_synthese(wb, analyse, certaines, a_controler, arbitrages, contexte) -> None:
    ws = wb.create_sheet("Synthèse", 0)
    ws.append(["Doublons fournisseurs — synthèse"])
    ws["A1"].font = Font(bold=True, size=14, color=COULEUR_ENTETE)
    lignes = [
        ("Date du snapshot analysé", contexte.get("date_snapshot", "")),
        ("Périmètre", contexte.get("libelle_perimetre", "")),
        ("", ""),
        ("Familles de doublons actionnables", len(certaines) + len(a_controler)),
        ("  dont fusions certaines", len(certaines)),
        ("  dont fusions à contrôler", len(a_controler)),
        ("Arbitrages inter-tenants (aucune instruction émise)", len(arbitrages)),
        ("Groupes de restructuration (points de vente, à rattacher, pas à fusionner)", len(analyse["restructurations"])),
        ("Candidats de rattrapage à qualifier", len(analyse["candidats_rattrapage"])),
        ("Fiches en quarantaine d'identité", len(analyse["quarantaine"])),
    ]
    for label, valeur in lignes:
        ws.append([label, valeur])
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 30


def _feuille_fusions(wb, titre: str, familles, avec_controles: bool) -> None:
    entetes = [
        "SIRET",
        "Raison sociale (fiche à conserver)",
        "Fiche à conserver (ID)",
        "Fiches à fusionner (ID)",
        "Cause",
        "Client(s) concerné(s)",
    ]
    if avec_controles:
        entetes.append("Contrôles requis")
    entetes += COLONNES_SUIVI
    ws = _feuille(wb, titre, entetes)
    for fam in familles:
        maitre = next(m for m in fam.membres if m.id == fam.fiche_a_conserver_id)
        ligne = [
            fam.siret,
            maitre.raison_sociale,
            maitre.id,
            ", ".join(fam.fiches_a_fusionner_ids),
            LIBELLES_CAUSES.get(fam.cause, ""),
            _clients_concernes(fam.membres),
        ]
        if avec_controles:
            ligne.append(" ; ".join(LIBELLES_SIGNAUX[s] for s in fam.signaux))
        ligne += ["", "", "", ""]
        ws.append(ligne)
    _ajuster_largeurs(ws, [16, 34, 14, 24, 40, 26] + ([44] if avec_controles else []) + [12, 14, 16, 30])


def _feuille_arbitrages(wb, familles) -> None:
    ws = _feuille(
        wb,
        "Arbitrages inter-tenants",
        ["SIRET", "Raison sociale", "Clients concernés", "Fiches (ID)", "Lecture"] + COLONNES_SUIVI,
    )
    for fam in familles:
        ws.append(
            [
                fam.siret,
                fam.membres[0].raison_sociale,
                _clients_concernes(fam.membres),
                ", ".join(m.id for m in fam.membres),
                "Décision de gouvernance — pas un défaut de qualité, aucune instruction émise.",
                "",
                "",
                "",
                "",
            ]
        )
    _ajuster_largeurs(ws, [16, 34, 26, 24, 48, 12, 14, 16, 30])


def _feuille_restructurations(wb, groupes) -> None:
    ws = _feuille(
        wb,
        "Restructurations",
        ["SIREN", "Nombre d'établissements", "Fiches (ID / SIRET / ville)", "Lecture"] + COLONNES_SUIVI,
    )
    for g in groupes:
        detail = "; ".join(f"{m.id} / {m.identite.siret} / {m.ville}" for m in g.membres)
        ws.append(
            [
                g.siren,
                g.nb_etablissements,
                detail,
                "Points de vente d'une même entreprise — rattacher, ne pas fusionner.",
                "",
                "",
                "",
                "",
            ]
        )
    _ajuster_largeurs(ws, [14, 20, 60, 46, 12, 14, 16, 30])


def _feuille_candidats(wb, candidats) -> None:
    ws = _feuille(wb, "Candidats rattrapage", ["Indice", "Type d'indice", "Fiches (ID)"] + COLONNES_SUIVI)
    for c in candidats:
        ws.append([c.cle_indice, c.type_indice, ", ".join(m.id for m in c.membres), "", "", "", ""])
    _ajuster_largeurs(ws, [30, 22, 34, 12, 14, 16, 30])


def _feuille_quarantaine(wb, lignes) -> None:
    ws = _feuille(
        wb, "Quarantaine identité", ["ID fiche", "Raison sociale", "Motif", "Correction suggérée"] + COLONNES_SUIVI
    )
    for q in lignes:
        ws.append([q.fournisseur.id, q.fournisseur.raison_sociale, q.motif, q.correction_suggeree or "", "", "", "", ""])
    _ajuster_largeurs(ws, [12, 34, 60, 30, 12, 14, 16, 30])


def construire_synthese_json(analyse: dict, contexte: dict) -> dict:
    familles = analyse["familles"]
    certaines = [f for f in familles if f.population == Population.DOUBLON_ACTIONNABLE and f.palier == Palier.CERTAINE]
    a_controler = [f for f in familles if f.population == Population.DOUBLON_ACTIONNABLE and f.palier == Palier.A_CONTROLER]
    arbitrages = [f for f in familles if f.population == Population.ARBITRAGE_INTER_TENANT]
    return {
        "date_snapshot": contexte.get("date_snapshot"),
        "perimetre": contexte.get("libelle_perimetre"),
        "doublons_actionnables": {
            "fusions_certaines": len(certaines),
            "fusions_a_controler": len(a_controler),
        },
        "arbitrages_inter_tenants": len(arbitrages),
        "restructurations": len(analyse["restructurations"]),
        "candidats_rattrapage": len(analyse["candidats_rattrapage"]),
        "quarantaine_identite": len(analyse["quarantaine"]),
    }


def ecrire_synthese_json(analyse: dict, chemin: Path, contexte: dict) -> Path:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_text(
        json.dumps(construire_synthese_json(analyse, contexte), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return chemin
